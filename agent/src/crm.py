import os
import json
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from config.settings import CRM_FILE


# Colour scheme for lead status
STATUS_FILLS = {
    "hot": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
    "warm": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
    "cold": PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid"),
    "client": PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid"),
}

CONTACT_COLUMNS = [
    "Name", "Phone", "WeChat ID", "Lead Status", "Property Type",
    "Budget Min", "Budget Max", "Preferred Locations", "Timeline",
    "Notes", "Last Contact", "Suggested Follow-up", "Follow-up Date",
    "Conversation History"
]

FOLLOWUP_COLUMNS = [
    "Priority", "Contact Name", "Action", "Due Date", "Lead Status", "Notes"
]


def init_crm():
    """Create a new CRM spreadsheet if one doesn't exist."""
    if os.path.exists(CRM_FILE):
        return

    wb = Workbook()

    # Contacts sheet
    ws_contacts = wb.active
    ws_contacts.title = "Contacts"
    _write_header(ws_contacts, CONTACT_COLUMNS)

    # Follow-ups sheet
    ws_followups = wb.create_sheet("Follow-ups")
    _write_header(ws_followups, FOLLOWUP_COLUMNS)

    # Dashboard sheet
    ws_dashboard = wb.create_sheet("Dashboard")
    ws_dashboard["A1"] = "WeChat CRM Dashboard"
    ws_dashboard["A1"].font = Font(size=16, bold=True)
    ws_dashboard["A3"] = "Total Contacts"
    ws_dashboard["B3"] = 0
    ws_dashboard["A4"] = "Hot Leads"
    ws_dashboard["B4"] = 0
    ws_dashboard["A5"] = "Warm Leads"
    ws_dashboard["B5"] = 0
    ws_dashboard["A6"] = "Cold Leads"
    ws_dashboard["B6"] = 0
    ws_dashboard["A7"] = "Clients"
    ws_dashboard["B7"] = 0
    ws_dashboard["A8"] = "Follow-ups Due Today"
    ws_dashboard["B8"] = 0
    ws_dashboard["A10"] = "Last Updated"
    ws_dashboard["B10"] = datetime.now().strftime("%Y-%m-%d %H:%M")

    wb.save(CRM_FILE)
    print(f"[CRM] Created new CRM at {CRM_FILE}")


def _write_header(ws, columns):
    """Write styled header row."""
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
    thin_border = Border(
        bottom=Side(style="thin", color="333333")
    )
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = max(len(col_name) + 4, 15)


def load_contacts() -> list[dict]:
    """Load all contacts from the CRM spreadsheet."""
    if not os.path.exists(CRM_FILE):
        init_crm()
        return []

    wb = load_workbook(CRM_FILE)
    ws = wb["Contacts"]

    contacts = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        values = [cell.value for cell in row]
        if not values[0]:  # skip empty rows
            continue
        contact = {}
        for i, col_name in enumerate(CONTACT_COLUMNS):
            key = col_name.lower().replace(" ", "_").replace("-", "_")
            contact[key] = values[i] if i < len(values) else None
        contacts.append(contact)

    return contacts


def update_contacts(extracted_contacts: list[dict]) -> list[str]:
    """
    Merge extracted contacts into the CRM.
    Returns a list of update descriptions for the briefing email.
    """
    init_crm()
    wb = load_workbook(CRM_FILE)
    ws = wb["Contacts"]

    existing = {}
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0]:
            existing[row[0].strip().lower()] = row_idx

    updates = []
    today = datetime.now().strftime("%Y-%m-%d")

    for contact in extracted_contacts:
        name = contact.get("name", "").strip()
        if not name:
            continue

        name_key = name.lower()
        locations = contact.get("preferred_locations", [])
        if isinstance(locations, list):
            locations = ", ".join(locations)

        if name_key in existing:
            # Update existing contact — never downgrade lead status
            row_idx = existing[name_key]
            old_status = (ws.cell(row=row_idx, column=4).value or "cold").lower()
            new_status = contact.get("lead_status", old_status).lower()

            status_rank = {"cold": 0, "warm": 1, "hot": 2, "client": 3}
            if status_rank.get(new_status, 0) > status_rank.get(old_status, 0):
                ws.cell(row=row_idx, column=4, value=new_status)
                _apply_status_fill(ws, row_idx, new_status)
                updates.append(f"{name}: upgraded from {old_status} to {new_status}")

            # Update fields if new data is available
            _update_if_new(ws, row_idx, 2, contact.get("phone"))
            _update_if_new(ws, row_idx, 3, contact.get("wechat_id"))
            _update_if_new(ws, row_idx, 5, contact.get("property_type"))
            _update_if_new(ws, row_idx, 6, contact.get("budget_min"))
            _update_if_new(ws, row_idx, 7, contact.get("budget_max"))
            _update_if_new(ws, row_idx, 8, locations)
            _update_if_new(ws, row_idx, 9, contact.get("timeline"))

            # Always update these
            ws.cell(row=row_idx, column=11, value=today)
            if contact.get("notes"):
                old_notes = ws.cell(row=row_idx, column=10).value or ""
                ws.cell(row=row_idx, column=10, value=f"{old_notes}\n[{today}] {contact['notes']}".strip())
            if contact.get("suggested_followup"):
                ws.cell(row=row_idx, column=12, value=contact["suggested_followup"])
            if contact.get("followup_date"):
                ws.cell(row=row_idx, column=13, value=contact["followup_date"])

            updates.append(f"{name}: updated contact info")
        else:
            # New contact
            new_row = ws.max_row + 1
            status = contact.get("lead_status", "cold").lower()
            ws.cell(row=new_row, column=1, value=name)
            ws.cell(row=new_row, column=2, value=contact.get("phone"))
            ws.cell(row=new_row, column=3, value=contact.get("wechat_id"))
            ws.cell(row=new_row, column=4, value=status)
            ws.cell(row=new_row, column=5, value=contact.get("property_type"))
            ws.cell(row=new_row, column=6, value=contact.get("budget_min"))
            ws.cell(row=new_row, column=7, value=contact.get("budget_max"))
            ws.cell(row=new_row, column=8, value=locations)
            ws.cell(row=new_row, column=9, value=contact.get("timeline"))
            ws.cell(row=new_row, column=10, value=contact.get("notes", ""))
            ws.cell(row=new_row, column=11, value=today)
            ws.cell(row=new_row, column=12, value=contact.get("suggested_followup"))
            ws.cell(row=new_row, column=13, value=contact.get("followup_date"))
            _apply_status_fill(ws, new_row, status)
            updates.append(f"New contact: {name} ({status})")

    # Rebuild follow-ups sheet
    _rebuild_followups(wb)

    # Update dashboard
    _update_dashboard(wb)

    wb.save(CRM_FILE)
    return updates


def _update_if_new(ws, row_idx, col_idx, value):
    """Only update a cell if the new value is non-empty and the old is empty."""
    if value and not ws.cell(row=row_idx, column=col_idx).value:
        ws.cell(row=row_idx, column=col_idx, value=value)


def _apply_status_fill(ws, row_idx, status):
    """Apply colour fill to the status cell."""
    fill = STATUS_FILLS.get(status)
    if fill:
        ws.cell(row=row_idx, column=4).fill = fill
        if status == "client":
            ws.cell(row=row_idx, column=4).font = Font(color="FFFFFF", bold=True)


def _rebuild_followups(wb):
    """Rebuild the Follow-ups sheet from Contacts data."""
    ws_contacts = wb["Contacts"]
    ws_followups = wb["Follow-ups"]

    # Clear existing data (keep header)
    for row in ws_followups.iter_rows(min_row=2):
        for cell in row:
            cell.value = None

    today = datetime.now().date()
    followups = []

    for row in ws_contacts.iter_rows(min_row=2, values_only=True):
        name = row[0]
        if not name or not row[12]:  # no follow-up date
            continue

        try:
            followup_date = datetime.strptime(str(row[12])[:10], "%Y-%m-%d").date()
        except (ValueError, TypeError):
            continue

        days_until = (followup_date - today).days
        if days_until <= 0:
            priority = "TODAY"
        elif days_until <= 7:
            priority = "This week"
        elif days_until <= 30:
            priority = "This month"
        else:
            priority = "Later"

        followups.append({
            "priority": priority,
            "name": name,
            "action": row[11] or "Follow up",
            "date": str(row[12])[:10],
            "status": row[3] or "cold",
            "notes": (row[9] or "")[:100],
            "sort_key": days_until
        })

    # Sort by urgency
    priority_order = {"TODAY": 0, "This week": 1, "This month": 2, "Later": 3}
    followups.sort(key=lambda x: (priority_order.get(x["priority"], 4), x["sort_key"]))

    priority_fills = {
        "TODAY": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
        "This week": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
        "This month": PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid"),
    }

    for idx, fu in enumerate(followups, start=2):
        ws_followups.cell(row=idx, column=1, value=fu["priority"])
        ws_followups.cell(row=idx, column=2, value=fu["name"])
        ws_followups.cell(row=idx, column=3, value=fu["action"])
        ws_followups.cell(row=idx, column=4, value=fu["date"])
        ws_followups.cell(row=idx, column=5, value=fu["status"])
        ws_followups.cell(row=idx, column=6, value=fu["notes"])

        fill = priority_fills.get(fu["priority"])
        if fill:
            ws_followups.cell(row=idx, column=1).fill = fill


def _update_dashboard(wb):
    """Update the Dashboard sheet with current stats."""
    ws_contacts = wb["Contacts"]
    ws_dashboard = wb["Dashboard"]

    counts = {"hot": 0, "warm": 0, "cold": 0, "client": 0}
    total = 0
    today = datetime.now().strftime("%Y-%m-%d")
    followups_today = 0

    for row in ws_contacts.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        total += 1
        status = (row[3] or "cold").lower()
        counts[status] = counts.get(status, 0) + 1
        if row[12] and str(row[12])[:10] <= today:
            followups_today += 1

    ws_dashboard["B3"] = total
    ws_dashboard["B4"] = counts["hot"]
    ws_dashboard["B5"] = counts["warm"]
    ws_dashboard["B6"] = counts["cold"]
    ws_dashboard["B7"] = counts["client"]
    ws_dashboard["B8"] = followups_today
    ws_dashboard["B10"] = datetime.now().strftime("%Y-%m-%d %H:%M")


def export_crm_json() -> dict:
    """Export CRM data as JSON for the web dashboard."""
    contacts = load_contacts()
    today = datetime.now().date()

    followups = []
    for c in contacts:
        if c.get("follow_up_date"):
            try:
                fu_date = datetime.strptime(str(c["follow_up_date"])[:10], "%Y-%m-%d").date()
                days_until = (fu_date - today).days
                if days_until <= 0:
                    priority = "today"
                elif days_until <= 7:
                    priority = "this_week"
                elif days_until <= 30:
                    priority = "this_month"
                else:
                    priority = "later"
                followups.append({
                    "contact_name": c["name"],
                    "action": c.get("suggested_follow_up", "Follow up"),
                    "due_date": str(c["follow_up_date"])[:10],
                    "priority": priority,
                    "lead_status": c.get("lead_status", "cold"),
                })
            except (ValueError, TypeError):
                pass

    counts = {"hot": 0, "warm": 0, "cold": 0, "client": 0}
    for c in contacts:
        status = (c.get("lead_status") or "cold").lower()
        counts[status] = counts.get(status, 0) + 1

    return {
        "updated_at": datetime.now().isoformat(),
        "contacts": contacts,
        "followups": followups,
        "stats": {
            "total": len(contacts),
            **counts,
            "followups_due_today": sum(1 for f in followups if f["priority"] == "today"),
        }
    }
